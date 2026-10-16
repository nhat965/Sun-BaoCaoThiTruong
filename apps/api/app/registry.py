import csv
import json
from collections.abc import Iterable
from datetime import UTC, datetime
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from app.database import SessionLocal, init_db
from app.repositories import SecuritiesRepository
from app.schemas import SecuritiesImportResult, Security

ALLOWED_EXCHANGES = {"HOSE", "HNX", "UPCOM"}
ALLOWED_STATUSES = {"active", "inactive", "delisted", "suspended"}


def parse_security(row: dict[str, str], default_source: str) -> Security:
    symbol = row.get("symbol", "").strip().upper()
    exchange = row.get("exchange", "").strip().upper()
    status = row.get("listing_status", "active").strip().lower()
    if not symbol.isalnum() or not 1 <= len(symbol) <= 16:
        raise ValueError("symbol must be 1-16 alphanumeric characters")
    if exchange not in ALLOWED_EXCHANGES:
        raise ValueError(f"exchange must be one of {sorted(ALLOWED_EXCHANGES)}")
    if status not in ALLOWED_STATUSES:
        raise ValueError(f"listing_status must be one of {sorted(ALLOWED_STATUSES)}")
    name = row.get("name", "").strip()
    if not name:
        raise ValueError("name is required")
    provider_ids = {}
    raw_provider_ids = row.get("provider_ids", "").strip()
    if raw_provider_ids:
        parsed = json.loads(raw_provider_ids)
        if not isinstance(parsed, dict) or not all(
            isinstance(key, str) and isinstance(value, str)
            for key, value in parsed.items()
        ):
            raise ValueError("provider_ids must be a JSON object of string identifiers")
        provider_ids.update({key.strip(): value.strip() for key, value in parsed.items()})
    provider_symbol = row.get("provider_symbol", "").strip()
    if provider_symbol:
        provider_ids["yahoo"] = provider_symbol
    for key, value in row.items():
        if key.startswith("provider_") and key not in {"provider_ids", "provider_symbol"}:
            provider_name = key.removeprefix("provider_").strip()
            if provider_name and value.strip():
                provider_ids[provider_name] = value.strip()
    return Security(
        symbol=symbol,
        name=name,
        aliases=[item.strip() for item in row.get("aliases", "").split("|") if item.strip()],
        exchange=exchange,
        industry=row.get("industry") or None,
        sector=row.get("sector") or None,
        listing_status=status,
        security_type=row.get("security_type", "stock") or "stock",
        provider_ids=provider_ids,
        source_name=row.get("source_name") or default_source,
        source_url=row.get("source_url") or None,
        last_updated=datetime.now(UTC),
    )


def import_security_rows(
    rows: Iterable[tuple[int, dict[str, str]]], source_name: str
) -> SecuritiesImportResult:
    # Keep the file importer usable as a standalone operator utility as well as
    # an HTTP dependency.  The API startup normally initializes the schema, but
    # direct CSV/XLSX imports must not depend on request ordering or a prior app
    # bootstrapping call.
    init_db()
    inserted = updated = rejected = 0
    errors: list[str] = []
    seen: set[str] = set()
    with SessionLocal() as session:
        repository = SecuritiesRepository(session)
        for line_number, row in rows:
            try:
                symbol = row.get("symbol", "").strip().upper()
                if symbol in seen:
                    raise ValueError("duplicate symbol in import")
                seen.add(symbol)
                _, was_inserted = repository.upsert(parse_security(row, source_name))
                inserted += int(was_inserted)
                updated += int(not was_inserted)
            except (ValueError, ValidationError) as exc:
                rejected += 1
                errors.append(f"line {line_number}: {exc}")
        session.commit()
    return SecuritiesImportResult(
        inserted=inserted, updated=updated, rejected=rejected, errors=errors[:100]
    )


def import_securities_csv(content: str, source_name: str) -> SecuritiesImportResult:
    rows = enumerate(csv.DictReader(StringIO(content)), start=2)
    return import_security_rows(rows, source_name)


def import_securities_xlsx(content: bytes, source_name: str) -> SecuritiesImportResult:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise ValueError("XLSX workbook is invalid or unreadable") from exc
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    try:
        header = [str(value or "").strip() for value in next(values)]
    except StopIteration as exc:
        raise ValueError("XLSX workbook is empty") from exc
    required = {"symbol", "name", "exchange"}
    if not required.issubset(header):
        raise ValueError(f"XLSX requires columns: {', '.join(sorted(required))}")
    rows = (
        (
            line_number,
            {
                key: "" if value is None else str(value).strip()
                for key, value in zip(header, row, strict=False)
            },
        )
        for line_number, row in enumerate(values, start=2)
    )
    return import_security_rows(rows, source_name)


def seed_registry_if_empty() -> SecuritiesImportResult | None:
    with SessionLocal() as session:
        if SecuritiesRepository(session).search("", limit=1, include_inactive=True):
            return None
    seed_path = Path(__file__).with_name("data") / "securities_bootstrap.csv"
    return import_securities_csv(seed_path.read_text(encoding="utf-8"), "bootstrap-manual")
